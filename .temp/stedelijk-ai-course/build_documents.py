from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\MData\Kennis\Courses\Stedelijk Gymnasium Leiden - AI voor docenten")
INK = "17222C"
BURGUNDY = "6B1F2A"
GOLD = "D6A756"
LIGHT = "F5F0E8"
PALE = "E8EEF5"
MUTED = "667085"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_inches):
    width = int(width_inches * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    widths_dxa = [round(width * 1440) for width in widths]
    widths_dxa[-1] += 9360 - sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def style_doc(doc, running_title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in [
        ("Title", 28, 0, 12, BURGUNDY),
        ("Heading 1", 16, 18, 10, BURGUNDY),
        ("Heading 2", 13, 14, 7, BURGUNDY),
        ("Heading 3", 12, 10, 5, INK),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ["List Bullet", "List Number"]:
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = running_title
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    header.paragraph_format.space_after = Pt(3)
    p_pr = header._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), GOLD)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    run = sub.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_callout(doc, heading, text, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(heading)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BURGUNDY)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, PALE)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if widths:
                set_cell_width(cells[i], widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(value))
    set_table_geometry(table, widths or [6.5 / len(headers)] * len(headers))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_page_break()


def build_workbook():
    doc = Document()
    style_doc(doc, "AI voor docenten - deelnemerswerkboek")
    add_title(doc, "AI voor docenten", "Deelnemerswerkboek | Stedelijk Gymnasium Leiden | 2 uur")
    add_callout(
        doc,
        "Doel",
        "Aan het einde heb je één bruikbare AI-workflow, een verbeterde toets- of lesopdracht en een veilige werkwijze die je morgen kunt toepassen.",
    )

    doc.add_heading("Mijn beginsituatie", level=1)
    add_table(
        doc,
        ["Vraag", "Mijn antwoord"],
        [
            ("Welke taak kost mij onnodig veel tijd?", ""),
            ("Waar gebruik ik AI al voor?", ""),
            ("Waar wil ik AI juist niet voor gebruiken?", ""),
            ("Wat wil ik na 2 uur kunnen?", ""),
        ],
        [2.3, 4.2],
    )

    doc.add_heading("De basis in één minuut", level=1)
    bullet(doc, "Generatieve AI voorspelt een passend vervolg; het controleert waarheid niet automatisch.")
    bullet(doc, "Een goede prompt maakt doel, context, bronnen, criteria, format en controle expliciet.")
    bullet(doc, "AI-output is een concept. Jij controleert vakinhoud, didactiek, privacy, bias en bronnen.")
    bullet(doc, "Geen persoonsgegevens of vertrouwelijke informatie in een niet-goedgekeurde tool.")
    bullet(doc, "Beoordeling en besluiten blijven menselijk.")

    doc.add_heading("Mijn vaste promptcanvas", level=1)
    add_table(
        doc,
        ["Onderdeel", "Invulling"],
        [
            ("Doel", "Ik wil..."),
            ("Context", "Vak, klas, niveau, leerdoel, voorkennis..."),
            ("Bronnen", "Gebruik alleen... / meld ontbrekende informatie."),
            ("Criteria", "Het resultaat moet..."),
            ("Format", "Lever op als..."),
            ("Controle", "Noem aannames, onzekerheden en mijn controlepunten."),
        ],
        [1.4, 5.1],
    )

    page_break(doc)
    doc.add_heading("Hands-on 1 - bouw een betere toets", level=1)
    doc.add_paragraph("Werk in duo's. Kies één echt leerdoel, maar gebruik geen herkenbare leerlinggegevens.")
    numbered(doc, "Schrijf het leerdoel meetbaar op.")
    numbered(doc, "Laat AI een mini-toetsmatrijs met 3-5 vragen voorstellen.")
    numbered(doc, "Laat AI ook een antwoordmodel en puntentoekenning maken.")
    numbered(doc, "Valideer met de kwaliteitscheck hieronder.")
    numbered(doc, "Verbeter de prompt één keer en bewaar de betere versie.")

    doc.add_heading("Kwaliteitscheck toets", level=2)
    add_table(
        doc,
        ["Controle", "Ja / aanpassen", "Notitie"],
        [
            ("Elke vraag meet het bedoelde leerdoel.", "", ""),
            ("Vraagtaal past bij klas en vak.", "", ""),
            ("Er is één duidelijke opdracht per vraag.", "", ""),
            ("Antwoordmodel accepteert verdedigbare varianten.", "", ""),
            ("Punten weerspiegelen denkstappen en moeilijkheid.", "", ""),
            ("Feiten, bronnen en berekeningen zijn gecontroleerd.", "", ""),
            ("De toets is haalbaar binnen de beschikbare tijd.", "", ""),
        ],
        [3.3, 1.2, 2.0],
    )

    doc.add_heading("Prompt voor deze opdracht", level=2)
    doc.add_paragraph(
        "Ontwerp een mini-toets voor [vak, klas] bij dit leerdoel: [leerdoel]. "
        "Maak eerst een toetsmatrijs en daarna 3-5 vragen met antwoordmodel en punten. "
        "Gebruik uitsluitend deze bron: [bron]. Controleer ambiguïteit, inhoudelijke juistheid, "
        "taalniveau en haalbaarheid. Sluit af met onzekerheden die ik als docent moet controleren."
    )

    doc.add_heading("Mijn verbeteringen", level=2)
    add_table(
        doc,
        ["Wat AI voorstelde", "Wat ik veranderde", "Waarom"],
        [("", "", ""), ("", "", ""), ("", "", "")],
        [2.15, 2.15, 2.2],
    )

    page_break(doc)
    doc.add_heading("Hands-on 2 - AI als tweede lezer", level=1)
    doc.add_paragraph(
        "Gebruik een fictief antwoord of het oefenantwoord hieronder. Laat AI bewijs uit de tekst koppelen aan de rubric. "
        "AI geeft geen eindcijfer."
    )
    doc.add_heading("Oefenopdracht", level=2)
    doc.add_paragraph(
        "Leg in 150-200 woorden uit waarom de val van de Romeinse Republiek niet door één oorzaak kan worden verklaard. "
        "Gebruik politieke, sociale en militaire factoren en leg minimaal één wisselwerking uit."
    )
    doc.add_heading("Fictief leerlingantwoord", level=2)
    doc.add_paragraph(
        "De Romeinse Republiek viel vooral omdat generaals steeds machtiger werden. Soldaten waren loyaal aan hun generaal "
        "omdat die hen betaalde en land kon geven. Daardoor konden generaals zoals Caesar hun leger gebruiken in politieke "
        "conflicten. Tegelijk was er veel ongelijkheid. Kleine boeren verloren grond en trokken naar Rome, waar politici "
        "hun steun probeerden te winnen. De Senaat verdedigde vooral de belangen van de elite en reageerde vaak te laat op "
        "hervormingen. De veroveringen maakten deze problemen groter: ze brachten rijkdom en slaven, maar ook machtige legers. "
        "De militaire, sociale en politieke oorzaken versterkten elkaar dus. Caesar was belangrijk, maar hij was eerder een "
        "gevolg van langere spanningen dan de enige oorzaak."
    )
    doc.add_heading("Compacte rubric", level=2)
    add_table(
        doc,
        ["Criterium", "Sterk", "Basis", "Onvoldoende"],
        [
            ("Inhoud", "Drie relevante factoren correct uitgelegd.", "Twee factoren grotendeels correct.", "Eén factor of meerdere fouten."),
            ("Samenhang", "Wisselwerking duidelijk en logisch.", "Verband genoemd maar beperkt uitgelegd.", "Opsomming zonder verband."),
            ("Onderbouwing", "Passende voorbeelden ondersteunen redenering.", "Minimaal één passend voorbeeld.", "Geen of onjuist voorbeeld."),
        ],
        [1.2, 1.8, 1.8, 1.7],
    )
    doc.add_heading("Vergelijking", level=2)
    add_table(
        doc,
        ["Vraag", "Mijn oordeel", "AI-voorstel", "Eindbesluit + reden"],
        [
            ("Welk tekstbewijs past bij elk criterium?", "", "", ""),
            ("Wat is de belangrijkste feedforward-tip?", "", "", ""),
            ("Waar was AI te stellig of te vaag?", "", "", ""),
        ],
        [2.0, 1.4, 1.4, 1.7],
    )

    page_break(doc)
    doc.add_heading("AI-signalen: van vermoeden naar zorgvuldig onderzoek", level=1)
    add_callout(
        doc,
        "Geen detectorvonnis",
        "Een detectorpercentage of 'AI-achtige stijl' bewijst niets. Gebruik signalen alleen als aanleiding voor een neutrale controle van proces, bronnen en begrip.",
        "FFF1F0",
    )
    doc.add_heading("Bewijsladder", level=2)
    add_table(
        doc,
        ["Sterkte", "Voorbeeld", "Betekenis"],
        [
            ("Zwak", "Stijlbreuk, glad taalgebruik, detectorflag", "Alleen aanleiding om beter te kijken."),
            ("Middel", "Bronnen bestaan niet, tekst past niet bij opdracht, geen tussenwerk", "Bespreek proces en vraag toelichting."),
            ("Sterker", "Versiegeschiedenis, promptlog of erkenning", "Beoordeel volgens vooraf bekende regels."),
            ("Beslissend", "Nooit één technisch signaal op zichzelf", "Menselijk onderzoek en hoor-wederhoor blijven nodig."),
        ],
        [1.0, 2.7, 2.8],
    )
    doc.add_heading("Neutrale gespreksvragen", level=2)
    bullet(doc, "Kun je uitleggen hoe je van bron naar deze conclusie bent gekomen?")
    bullet(doc, "Welke tussenversies of aantekeningen heb je?")
    bullet(doc, "Waarom koos je juist deze voorbeelden en formulering?")
    bullet(doc, "Welke hulpmiddelen heb je gebruikt, en waarvoor?")
    bullet(doc, "Wat zou je nu inhoudelijk veranderen?")

    doc.add_heading("Maak de opdracht procesrijker", level=2)
    add_table(
        doc,
        ["Huidige opdracht", "Procesbewijs", "Toegestaan AI-gebruik", "Controle"],
        [("", "", "", ""), ("", "", "", "")],
        [1.6, 1.7, 1.8, 1.4],
    )

    page_break(doc)
    doc.add_heading("Veiligheidskaart", level=1)
    add_table(
        doc,
        ["Zone", "Voorbeelden", "Actie"],
        [
            ("GROEN", "Brainstorm, fictieve voorbeelden, lesstructuur, quizvarianten", "Mag in goedgekeurde tool; controleer output."),
            ("ORANJE", "Geanonimiseerd werk, feedbackconcept, toetsmateriaal", "Alleen met doel, minimale data en menselijke eindcontrole."),
            ("ROOD", "Namen, cijfers, dossiers, medische info, centraal examenwerk, disciplinaire besluiten", "Niet invoeren; volg formele schoolprocedure."),
        ],
        [1.0, 3.1, 2.4],
    )

    doc.add_heading("Mijn persoonlijke workflow", level=1)
    add_table(
        doc,
        ["Taak", "AI doet", "Ik controleer", "Gegevensrisico", "Eerste testdatum"],
        [("", "", "", "", ""), ("", "", "", "", "")],
        [1.2, 1.5, 1.5, 1.2, 1.1],
    )
    add_callout(
        doc,
        "Mijn 48-uursactie",
        "Binnen 48 uur gebruik ik AI voor: ____________________. Ik controleer het resultaat op: ____________________. "
        "Ik deel mijn bevinding met: ____________________.",
    )

    doc.add_heading("AI-outputcheck: K.L.O.P.T.", level=1)
    bullet(doc, "K - Klopt de vakinhoud en zijn berekeningen juist?")
    bullet(doc, "L - Ligt het resultaat in lijn met leerdoel, niveau en schoolafspraken?")
    bullet(doc, "O - Onderbouwt de output claims met controleerbare bronnen?")
    bullet(doc, "P - Privacy: zijn gegevens fictief of afdoende geanonimiseerd?")
    bullet(doc, "T - Teken ik zelf voor de eindbeslissing en kan ik die uitleggen?")

    path = ROOT / "02-Docentenwerkboek.docx"
    doc.save(path)
    return path


def build_trainer_guide():
    doc = Document()
    style_doc(doc, "AI voor docenten - trainershandleiding")
    add_title(doc, "Trainershandleiding", "AI voor docenten | Stedelijk Gymnasium Leiden | 120 minuten")
    add_callout(
        doc,
        "Trainingsbelofte",
        "Docenten vertrekken niet alleen met inspiratie, maar met één geteste workflow, een promptbibliotheek, een veiligheidskaart en een concept-SOP.",
    )

    doc.add_heading("Succescriteria", level=1)
    bullet(doc, "Minimaal 80% van de deelnemers maakt een werkende prompt met expliciete controle-eisen.")
    bullet(doc, "Elke deelnemer verbetert één echte les-, toets- of feedbacktaak.")
    bullet(doc, "Deelnemers kunnen uitleggen waarom detectoren geen zelfstandig bewijs zijn.")
    bullet(doc, "Iedereen kiest één veilige toepassing voor de komende 48 uur.")

    doc.add_heading("Voorbereiding", level=1)
    bullet(doc, "Bevestig welke AI-tool(s) de school heeft goedgekeurd en welke accounts beschikbaar zijn.")
    bullet(doc, "Vraag deelnemers om een leerdoel, toetsvraag of rubric mee te nemen.")
    bullet(doc, "Open de presentatie en test de live demo met een fictieve context.")
    bullet(doc, "Deel het werkboek en de promptbibliotheek.")
    bullet(doc, "Stem vooraf af wie vragen over privacy, examenreglement en formeel beleid opvolgt.")

    doc.add_heading("Benodigdheden", level=1)
    bullet(doc, "Laptop per één of twee deelnemers, internet, beamer en timer.")
    bullet(doc, "Goedgekeurde AI-omgeving; geen verplicht persoonlijk leerlingaccount.")
    bullet(doc, "Optioneel: geprinte veiligheidskaart en SOP.")

    doc.add_heading("Draaiboek in één oogopslag", level=1)
    add_table(
        doc,
        ["Tijd", "Onderdeel", "Werkvorm", "Opbrengst"],
        [
            ("0-10", "Start en nulmeting", "Handen / duo-uitwisseling", "Persoonlijke leervraag"),
            ("10-25", "Werking, kansen, grenzen", "Uitleg + korte voorbeelden", "Realistisch mentaal model"),
            ("25-40", "Prompten en demo", "Live verbeteren", "Herbruikbaar promptcanvas"),
            ("40-65", "Hands-on 1", "Duo's: toets + model", "Gecontroleerd product"),
            ("65-70", "Pauze", "Staand reflecteren", "Energie + observatie"),
            ("70-90", "Feedback en nakijken", "Demo + vergelijking", "Tweede-lezerworkflow"),
            ("90-105", "AI-signalen", "Casus + gesprek", "Zorgvuldig handelingspad"),
            ("105-117", "Hands-on 2", "Eigen workflow + SOP", "48-uursactie"),
            ("117-120", "Afsluiting", "Commitment", "Eerstvolgende stap"),
        ],
        [0.8, 2.1, 1.8, 1.8],
    )

    page_break(doc)
    doc.add_heading("0-10 min - Start en nulmeting", level=1)
    doc.add_heading("Zeg", level=2)
    doc.add_paragraph(
        "\"Vandaag gaat niet over zoveel mogelijk AI gebruiken. Het gaat over weten waar AI je werk beter of sneller kan maken, "
        "waar het risico ontstaat en waar jij altijd zelf aan het stuur blijft.\""
    )
    doc.add_heading("Doe", level=2)
    numbered(doc, "Laat deelnemers één tijdrovende docenttaak opschrijven.")
    numbered(doc, "Laat hen per duo delen: kans, zorg en gewenste uitkomst.")
    numbered(doc, "Inventariseer met handen: nooit gebruikt / soms / wekelijks.")
    doc.add_heading("Let op", level=2)
    bullet(doc, "Normaliseer zowel enthousiasme als terughoudendheid.")
    bullet(doc, "Beloof geen exacte tijdwinst; die hangt af van taak, tool en controlelast.")

    doc.add_heading("10-25 min - Een bruikbaar mentaal model", level=1)
    doc.add_heading("Kernuitleg", level=2)
    bullet(doc, "Een taalmodel voorspelt plausibele tekst op basis van patronen.")
    bullet(doc, "Plausibel is niet hetzelfde als waar, passend of eerlijk.")
    bullet(doc, "Meer relevante context en controlecriteria verkleinen fouten, maar maken ze niet onmogelijk.")
    bullet(doc, "De beste rol is vaak: brainstormer, eerste versie, variantenmaker, structuurhulp of kritische tweede lezer.")
    doc.add_heading("Mini-demonstratie", level=2)
    doc.add_paragraph(
        "Vraag AI eerst: \"Maak een toets over de Romeinen.\" Bespreek waarom dit te vaag is. Gebruik daarna dezelfde taak met "
        "klas, leerdoel, bron, puntenverdeling, vraagtypen en kwaliteitscontrole. Laat deelnemers de verschillen benoemen."
    )

    doc.add_heading("25-40 min - Prompten en live demo", level=1)
    doc.add_heading("Demo-prompt", level=2)
    add_callout(
        doc,
        "Kopieer in de AI-tool",
        "Ontwerp een mini-toets voor geschiedenis, klas 2 gymnasium. Leerdoel: leerlingen leggen uit hoe militaire, sociale "
        "en politieke factoren elkaar versterkten bij de val van de Romeinse Republiek. Gebruik uitsluitend de bron die ik "
        "hieronder plak. Maak een toetsmatrijs, vier vragen voor 20 punten en een antwoordmodel. Verdeel over kennis, toepassing "
        "en redeneren. Controleer daarna ambiguïteit, alternatieve verdedigbare antwoorden, taalniveau en haalbaarheid in 25 minuten. "
        "Noem tot slot wat ik als docent nog moet verifiëren.",
    )
    doc.add_heading("Didactische vraag", level=2)
    doc.add_paragraph(
        "Laat de groep bepalen welk onderdeel van de prompt het meeste kwaliteitsverschil veroorzaakt. Het gewenste antwoord is "
        "niet één magisch format, maar de combinatie van helder leerdoel, beperkte bron en expliciete kwaliteitscontrole."
    )

    page_break(doc)
    doc.add_heading("40-65 min - Hands-on 1: toets plus antwoordmodel", level=1)
    doc.add_heading("Opdracht", level=2)
    numbered(doc, "Kies één leerdoel uit de eigen lespraktijk.")
    numbered(doc, "Maak met AI een toetsmatrijs en 3-5 vragen.")
    numbered(doc, "Laat AI een antwoordmodel maken.")
    numbered(doc, "Controleer met het werkblad; verbeter de prompt één keer.")
    numbered(doc, "Wissel met een duo en zoek één inhoudelijk of didactisch risico.")
    doc.add_heading("Trainer loopt rond en vraagt", level=2)
    bullet(doc, "Welk leerdoel wordt hier precies zichtbaar?")
    bullet(doc, "Welke bron begrenst de AI?")
    bullet(doc, "Welk antwoord zou een goede leerling kunnen geven dat nu ten onrechte wordt afgekeurd?")
    bullet(doc, "Wat moet jij vakinhoudelijk controleren?")
    doc.add_heading("Nabespreking", level=2)
    doc.add_paragraph(
        "Vraag drie duo's om alleen hun belangrijkste reparatie te delen. Leg nadruk op de menselijke kwaliteitswinst, niet op de eerste AI-output."
    )

    doc.add_heading("65-70 min - Korte pauze", level=1)
    doc.add_paragraph(
        "Laat deelnemers staand één zin afmaken: \"AI was nuttig toen..., maar ik moest ingrijpen bij...\" Verzamel twee voorbeelden."
    )

    doc.add_heading("70-90 min - Feedback en nakijken", level=1)
    doc.add_heading("Kader", level=2)
    bullet(doc, "Gebruik AI als tweede lezer: bewijs zoeken, rubric toepassen, feedbackopties voorstellen.")
    bullet(doc, "Geen autonoom cijfer, geen herkenbare leerlingdata, geen bulkbesluit zonder menselijke controle.")
    bullet(doc, "Bij centrale examens: geen AI bij het nakijken.")
    doc.add_heading("Demo", level=2)
    numbered(doc, "Toon het fictieve Romeinse antwoord uit het werkboek.")
    numbered(doc, "Laat AI per rubriccriterium tekstbewijs en een voorlopig niveau geven.")
    numbered(doc, "Vraag de groep waar AI te stellig, te mild of te oppervlakkig is.")
    numbered(doc, "Pas de rubric of prompt aan en laat één criterium opnieuw beoordelen.")
    doc.add_heading("Spreekzin", level=2)
    doc.add_paragraph(
        "\"Als AI en docent verschillen, wint niet automatisch de docent omdat die mens is, en zeker niet de AI omdat die snel is. "
        "We gaan terug naar leerdoel, rubric en bewijs uit het werk.\""
    )

    page_break(doc)
    doc.add_heading("90-105 min - AI-signalen en academische integriteit", level=1)
    doc.add_heading("Kernboodschap", level=2)
    add_callout(
        doc,
        "Detector = signaal, geen bewijs",
        "AI-detectors kunnen menselijke tekst als AI markeren en AI-tekst missen. Vooral bij meertalige schrijvers zijn "
        "problematische foutpositieven aangetoond. Gebruik een score nooit zelfstandig voor cijfer- of disciplinaire besluiten.",
        "FFF1F0",
    )
    doc.add_heading("Casus", level=2)
    doc.add_paragraph(
        "Een leerling levert plotseling een foutloze, formele tekst in. Een detector geeft 78% AI. De bronnenlijst bevat twee titels "
        "die niet vindbaar zijn. Wat doe je vandaag, wat leg je vast en wat doe je nadrukkelijk niet?"
    )
    doc.add_heading("Gewenst handelingspad", level=2)
    numbered(doc, "Controleer eerst opdrachtvoorwaarden, bronnen en feitelijke inconsistenties.")
    numbered(doc, "Vergelijk waar passend met beschikbaar proceswerk; interpreteer stijlverschil voorzichtig.")
    numbered(doc, "Voer een neutraal gesprek met vragen over keuzes, bronnen en tussenstappen.")
    numbered(doc, "Bied hoor en wederhoor en volg bestaande school- en examenregels.")
    numbered(doc, "Leg alleen relevante feiten vast; behandel detectoroutput niet als verdict.")
    doc.add_heading("Maak toetsen robuuster", level=2)
    bullet(doc, "Vraag tussenproducten of versiegeschiedenis.")
    bullet(doc, "Gebruik lokale, actuele of persoonlijke brondata waar dat didactisch past.")
    bullet(doc, "Voeg korte mondelinge verdediging of klassikale toepassing toe.")
    bullet(doc, "Maak toegestaan AI-gebruik vooraf expliciet en vraag een eenvoudig AI-logboek.")

    doc.add_heading("105-117 min - Hands-on 2: eigen workflow en SOP", level=1)
    numbered(doc, "Kies één terugkerende taak.")
    numbered(doc, "Bepaal de AI-rol: concept, varianten, structuur of tweede lezer.")
    numbered(doc, "Markeer de taak groen, oranje of rood.")
    numbered(doc, "Schrijf de menselijke controle en het stopmoment op.")
    numbered(doc, "Plan een kleine test binnen 48 uur.")
    doc.add_heading("Duo-check", level=2)
    doc.add_paragraph(
        "De collega probeert één fout, privacylek of onduidelijke verantwoordelijkheid in de workflow te vinden. De eigenaar past de workflow aan."
    )

    doc.add_heading("117-120 min - Afsluiting", level=1)
    doc.add_paragraph(
        "Laat iedereen de 48-uursactie hardop aan een buur noemen. Sluit af met: \"Begin klein. Geen persoonsgegevens. Controleer vakinhoud. "
        "Houd de beslissing menselijk. Deel wat werkt.\""
    )

    page_break(doc)
    doc.add_heading("Veelgestelde vragen", level=1)
    add_table(
        doc,
        ["Vraag", "Kort antwoord"],
        [
            ("Mag AI toetsen maken?", "Ja, als concept binnen goedgekeurde tools; docent valideert leerdoel, inhoud, moeilijkheid en beveiliging."),
            ("Mag AI nakijken?", "Als tweede lezer bij passende, geanonimiseerde situaties; niet autonoom en niet voor centrale examens."),
            ("Kan ik leerlingwerk uploaden?", "Alleen als de tool en verwerking formeel zijn goedgekeurd; standaard fictief of volledig geanonimiseerd."),
            ("Welke detector is het beste?", "Geen detector levert zelfstandig betrouwbaar bewijs. Richt je op proces, bronnen, begrip en hoor-wederhoor."),
            ("Moeten leerlingen AI gebruiken?", "Alleen wanneer dit didactisch nodig, veilig ingericht en passend bij leeftijd, opdracht en schoolbeleid is."),
            ("Wie is verantwoordelijk?", "De school voor kaders en systemen; de docent voor professioneel gebruik en eindbesluiten binnen die kaders."),
        ],
        [2.0, 4.5],
    )

    doc.add_heading("Plan B zonder werkende AI-tool", level=1)
    bullet(doc, "Gebruik screenshots of vooraf gegenereerde fictieve output.")
    bullet(doc, "Laat deelnemers prompts op papier verbeteren.")
    bullet(doc, "Laat duo's zelf 'AI spelen': één genereert snel, de ander controleert met K.L.O.P.T.")
    bullet(doc, "Behaal dezelfde leerdoelen; live toolgebruik is middel, geen doel.")

    doc.add_heading("Evaluatievragen", level=1)
    bullet(doc, "Wat ga je binnen 48 uur testen?")
    bullet(doc, "Welke rode lijn is voor jou duidelijker geworden?")
    bullet(doc, "Welke schoolafspraak ontbreekt nog?")
    bullet(doc, "Waar is vervolgtraining nodig: prompten, toetsontwerp, privacy of vakgerichte toepassingen?")

    path = ROOT / "03-Trainershandleiding.docx"
    doc.save(path)
    return path


def build_sop():
    doc = Document()
    style_doc(doc, "Concept AI-SOP en schoolrichtlijnen")
    add_title(doc, "AI-SOP en schoolrichtlijnen", "Concept 0.1 | Stedelijk Gymnasium Leiden | ter validatie")
    add_callout(
        doc,
        "Status",
        "Dit document is een direct bruikbaar werkconcept, geen vastgesteld schoolbeleid. Laat schoolleiding, privacyfunctionaris/FG en examencommissie het valideren.",
        "FFF4D6",
    )

    doc.add_heading("1. Doel en uitgangspunten", level=1)
    doc.add_paragraph(
        "Deze SOP helpt medewerkers generatieve AI doelgericht, veilig en uitlegbaar te gebruiken bij onderwijsvoorbereiding, "
        "feedback, toetsontwikkeling en administratieve concepttaken."
    )
    bullet(doc, "De mens blijft centraal en eindverantwoordelijk.")
    bullet(doc, "AI-gebruik moet een duidelijk onderwijs- of werkdoel dienen.")
    bullet(doc, "Gebruik alleen door school goedgekeurde tools en accounts.")
    bullet(doc, "Verwerk zo min mogelijk gegevens; standaard fictief of volledig geanonimiseerd.")
    bullet(doc, "AI-output wordt gecontroleerd op juistheid, bias, bronnen, niveau en privacy.")
    bullet(doc, "Leerlingen en medewerkers zijn transparant over relevant AI-gebruik.")

    doc.add_heading("2. Verkeerslichtmodel", level=1)
    add_table(
        doc,
        ["Zone", "Toepassing", "Voorwaarden"],
        [
            ("GROEN", "Brainstorm, lesopzet, fictieve voorbeelden, taalvarianten, quizvragen", "Goedgekeurde tool; geen persoonsgegevens; docent controleert."),
            ("ORANJE", "Geanonimiseerde tekstfeedback, toetsmateriaal, rubricconcept, foutpatronen", "Duidelijk doel; minimale data; extra vak- en privacycheck; menselijke eindbeslissing."),
            ("ROOD", "Namen, cijfers, leerlingdossiers, medische/ondersteuningsdata, vertrouwelijke communicatie, centraal examenwerk, disciplinaire of plaatsingsbesluiten", "Niet invoeren in generatieve AI; gebruik formele schoolprocedure."),
        ],
        [0.9, 2.7, 2.9],
    )

    doc.add_heading("3. De zesstappen-SOP", level=1)
    add_table(
        doc,
        ["Stap", "Actie", "Stop als..."],
        [
            ("1. Doel", "Formuleer het concrete onderwijs- of werkdoel.", "AI alleen 'omdat het kan' wordt gebruikt."),
            ("2. Tool", "Controleer of tool, account en functie zijn goedgekeurd.", "Goedkeuring of verwerkersafspraken onduidelijk zijn."),
            ("3. Data", "Verwijder namen en indirect herkenbare details; gebruik fictieve data.", "Een persoon redelijkerwijs herkenbaar blijft."),
            ("4. Prompt", "Geef context, bronnen, criteria, format en controleopdracht.", "Bronmateriaal vertrouwelijk of toetsbeveiligd is."),
            ("5. Controle", "Controleer K.L.O.P.T.; vergelijk met bron, rubric en vakkennis.", "Je de output niet zelfstandig kunt beoordelen."),
            ("6. Besluit", "Pas aan, leg relevant gebruik vast en neem zelf het besluit.", "AI feitelijk de beoordeling of beslissing overneemt."),
        ],
        [0.8, 3.2, 2.5],
    )

    doc.add_heading("4. K.L.O.P.T.-controle", level=1)
    bullet(doc, "K - Klopt de vakinhoud?")
    bullet(doc, "L - Ligt het in lijn met leerdoel, niveau en afspraken?")
    bullet(doc, "O - Onderbouwd met controleerbare bronnen en expliciete aannames?")
    bullet(doc, "P - Privacy beschermd en minimale gegevens gebruikt?")
    bullet(doc, "T - Tekent een mens voor de eindbeslissing en kan die deze uitleggen?")

    page_break(doc)
    doc.add_heading("5. Toetsontwikkeling", level=1)
    bullet(doc, "Start bij leerdoelen en toetsmatrijs; laat AI pas daarna vragen voorstellen.")
    bullet(doc, "Geef alleen toegestane bronnen en geen beveiligd bestaand toetsmateriaal in een publieke tool.")
    bullet(doc, "Controleer elke vraag op inhoud, ambiguïteit, taalbelasting, bias, moeilijkheid en haalbaarheid.")
    bullet(doc, "Maak antwoordmodel en puntentoekenning apart controleerbaar.")
    bullet(doc, "Bewaar definitieve toetsbestanden alleen in goedgekeurde schoolsystemen.")

    doc.add_heading("6. Feedback en nakijken", level=1)
    bullet(doc, "AI mag bewijs uit geanonimiseerd werk ordenen of feedbackopties voorstellen als de verwerking is toegestaan.")
    bullet(doc, "AI kent niet zelfstandig het eindcijfer toe en communiceert niet autonoom met leerlingen of ouders.")
    bullet(doc, "De docent vergelijkt AI-output met rubric, bronwerk en eigen vakinhoudelijk oordeel.")
    bullet(doc, "Bij verschil gaat men terug naar expliciete criteria en bewijs.")
    bullet(doc, "Gebruik geen AI bij het nakijken van centrale examens.")

    doc.add_heading("7. Mogelijk ongeoorloofd AI-gebruik", level=1)
    add_callout(
        doc,
        "Hoofdregel",
        "Een detectorflag, stijlbreuk of vermoeden is geen bewijs en mag niet zelfstandig leiden tot cijferverlaging of een disciplinaire maatregel.",
        "FFF1F0",
    )
    numbered(doc, "Controleer vooraf gecommuniceerde opdrachtvoorwaarden.")
    numbered(doc, "Controleer bronnen, feitelijke inconsistenties en beschikbaar procesbewijs.")
    numbered(doc, "Voer een neutraal gesprek over keuzes, bronnen, tussenstappen en hulpmiddelen.")
    numbered(doc, "Pas hoor en wederhoor toe en volg bestaande toets- en integriteitsprocedures.")
    numbered(doc, "Leg relevante feiten en het menselijke besluit vast; minimaliseer persoonsgegevens.")
    numbered(doc, "Gebruik de casus om opdrachtontwerp en voorlichting te verbeteren.")

    doc.add_heading("8. Transparantieverklaring voor leerlingen", level=1)
    add_callout(
        doc,
        "Voorbeeldtekst",
        "Ik heb [tool] gebruikt voor [doel]. Mijn belangrijkste prompts waren [kort]. Ik heb de output gecontroleerd met [bronnen/werkwijze]. "
        "Ik heb zelf [eigen bijdrage] uitgevoerd en kan alle inhoud en keuzes in dit werk uitleggen.",
    )

    doc.add_heading("9. Incidenten en twijfel", level=1)
    add_table(
        doc,
        ["Situatie", "Directe actie", "Escalatie"],
        [
            ("Per ongeluk persoonsgegevens ingevoerd", "Stop, verwijder waar mogelijk, maak geen nieuwe invoer.", "Meld volgens datalek/IBP-procedure."),
            ("Tool geeft schadelijke of discriminerende output", "Gebruik output niet; bewaar alleen wat nodig is voor onderzoek.", "Meld bij aangewezen AI/IBP-contact."),
            ("Onzeker of tool is toegestaan", "Gebruik de tool niet met schooldata.", "Vraag ICT/privacyfunctionaris."),
            ("Twijfel over beoordeling of examen", "Neem geen AI-besluit.", "Vraag vaksectie/examencommissie."),
        ],
        [1.7, 2.6, 2.2],
    )

    page_break(doc)
    doc.add_heading("10. Rollen en verantwoordelijkheden", level=1)
    add_table(
        doc,
        ["Rol", "Verantwoordelijkheid"],
        [
            ("Schoolleiding", "Visie, goedgekeurde toepassingen, middelen, training en toezicht."),
            ("Privacyfunctionaris/FG", "Privacyrisico's, DPIA, leveranciersafspraken en incidentadvies."),
            ("ICT/IBP", "Accounts, instellingen, toegang, logging, beveiliging en support."),
            ("Examencommissie", "Toets- en examenregels, toegestane inzet en casusafhandeling."),
            ("Vaksectie", "Vakinhoudelijke kwaliteitsnormen, rubrics en gedeelde voorbeelden."),
            ("Docent", "Doelgericht gebruik, dataminimalisatie, controle, transparantie en eindbesluit."),
            ("Leerling", "Volgt opdrachtregels, vermeldt relevant AI-gebruik en kan eigen werk uitleggen."),
        ],
        [1.7, 4.8],
    )

    doc.add_heading("11. Minimale schoolbesluiten voor invoering", level=1)
    add_table(
        doc,
        ["Besluit", "Eigenaar", "Status"],
        [
            ("Welke AI-tools en functies zijn goedgekeurd?", "Schoolleiding + ICT/IBP", "Te bepalen"),
            ("Welke gegevens mogen per tool worden verwerkt?", "Privacyfunctionaris/FG", "Te bepalen"),
            ("Welke regels gelden voor formatieve toetsen, schoolexamens en centrale examens?", "Examencommissie", "Te bepalen"),
            ("Welke transparantieverklaring gebruiken leerlingen?", "Onderwijsleiding + secties", "Te bepalen"),
            ("Waar worden incidenten en vragen gemeld?", "Schoolleiding", "Te bepalen"),
            ("Hoe wordt AI-geletterdheid jaarlijks onderhouden?", "Schoolleiding/HR", "Te bepalen"),
        ],
        [3.6, 1.9, 1.0],
    )

    doc.add_heading("12. Korte docentencheck voor publicatie of beoordeling", level=1)
    bullet(doc, "Ik gebruikte een goedgekeurde tool en een duidelijk doel.")
    bullet(doc, "Ik voerde geen onnodige of herkenbare persoonsgegevens in.")
    bullet(doc, "Ik controleerde vakinhoud, bronnen, niveau, bias en taal.")
    bullet(doc, "Ik kan uitleggen wat AI deed en wat ik zelf wijzigde.")
    bullet(doc, "Ik nam zelf de eindbeslissing.")

    doc.add_heading("13. Beheer", level=1)
    add_table(
        doc,
        ["Veld", "Waarde"],
        [
            ("Documenteigenaar", "Nog aan te wijzen"),
            ("Versie", "Concept 0.1"),
            ("Vastgesteld op", "Nog niet vastgesteld"),
            ("Herziening", "Minimaal jaarlijks en bij nieuwe tools/regelgeving"),
            ("Contactpunt", "Nog aan te wijzen"),
        ],
        [2.0, 4.5],
    )

    path = ROOT / "04-AI-SOP-en-schoolrichtlijnen.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    outputs = [build_workbook(), build_trainer_guide(), build_sop()]
    for output in outputs:
        print(output)
